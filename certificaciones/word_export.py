# certificaciones/word_export.py
"""Exporta un certificado a formato Word (.docx)."""
from io import BytesIO
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.shortcuts import get_object_or_404
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .models import Certificado


@login_required
def descargar_word(request, codigo):
    cert = get_object_or_404(Certificado, codigo=codigo)

    if not (request.user == cert.usuario or request.user.is_staff):
        return HttpResponse("No autorizado", status=403)

    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    title = doc.add_heading("CERTIFICADO DE FINALIZACION", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(cert.institucion)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(14)

    doc.add_paragraph()
    doc.add_paragraph("Se otorga el presente certificado a").alignment = WD_ALIGN_PARAGRAPH.CENTER

    nombre = doc.add_paragraph()
    nombre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = nombre.add_run(cert.usuario.get_full_name() or cert.usuario.username)
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0x0d, 0x6e, 0xfd)

    doc.add_paragraph("por haber completado satisfactoriamente el curso").alignment = WD_ALIGN_PARAGRAPH.CENTER

    curso_p = doc.add_paragraph()
    curso_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = curso_p.add_run('"' + cert.curso.titulo + '"')
    run.bold = True
    run.font.size = Pt(16)

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run("Nota final: " + str(cert.nota_final) + "/100  |  ")
    info.add_run("Duracion: " + str(cert.horas) + " horas")

    doc.add_paragraph()
    doc.add_paragraph("Emitido el " + cert.fecha_emision.strftime("%d de %B de %Y")).alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    doc.add_paragraph()

    firma = doc.add_paragraph()
    firma.alignment = WD_ALIGN_PARAGRAPH.CENTER
    firma.add_run("_" * 40 + "\n")
    admin_name = cert.firmado_por_admin.get_full_name() if cert.firmado_por_admin else "Director"
    firma.add_run(admin_name + "\n")
    firma.add_run(cert.cargo_admin or "Director General")

    doc.add_paragraph()
    firma2 = doc.add_paragraph()
    firma2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    firma2.add_run("_" * 40 + "\n")
    prof_name = cert.firmado_por_profesor.get_full_name() if cert.firmado_por_profesor else "Profesor"
    firma2.add_run(prof_name + "\n")
    firma2.add_run(cert.cargo_profesor or "Profesor Titular")

    doc.add_paragraph()
    codigo_p = doc.add_paragraph("Codigo de verificacion: " + str(cert.codigo))
    codigo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    codigo_p.runs[0].font.size = Pt(9)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    response["Content-Disposition"] = 'attachment; filename="certificado_' + str(cert.codigo)[:8] + '.docx"'
    return response